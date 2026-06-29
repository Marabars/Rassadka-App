from __future__ import annotations
import io
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import openpyxl.utils

from app.database import get_db

router = APIRouter()


@router.get("/seating/export")
async def export_seating(month: str):
    async with get_db() as db:
        rows = await (await db.execute(
            "SELECT employee_name, seat_id, status, date FROM seat_assignments"
            " WHERE month=? ORDER BY employee_name, date",
            (month,)
        )).fetchall()

    if not rows:
        raise HTTPException(404, "No data for this month")

    emp_data: dict[str, dict[str, str]] = {}
    all_dates: set[str] = set()
    for name, seat, status, date in rows:
        all_dates.add(date)
        if name not in emp_data:
            emp_data[name] = {}
        emp_data[name][date] = (seat or "") if status == "OFFICE" else ""

    sorted_dates = sorted(all_dates)
    sorted_names = sorted(emp_data.keys())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Рассадка {month}"

    hdr_fill  = PatternFill("solid", fgColor="1C1A35")
    hdr_font  = Font(bold=True, color="E5E3EB", name="Calibri")
    name_fill = PatternFill("solid", fgColor="13112A")
    name_font = Font(color="E5E3EB", name="Calibri")
    seat_font = Font(color="82D6CC", bold=True, name="Calibri")
    odd_fill  = PatternFill("solid", fgColor="13112A")
    even_fill = PatternFill("solid", fgColor="0F0D20")

    ws.cell(row=1, column=1, value="ФИО").fill = hdr_fill
    ws.cell(row=1, column=1).font = hdr_font
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left")
    ws.column_dimensions["A"].width = 36

    for col, date in enumerate(sorted_dates, start=2):
        dd_mm = date[8:10] + "." + date[5:7]
        cell = ws.cell(row=1, column=col, value=dd_mm)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 8

    for row_idx, name in enumerate(sorted_names, start=2):
        row_fill = odd_fill if row_idx % 2 == 0 else even_fill
        cell = ws.cell(row=row_idx, column=1, value=name)
        cell.fill = name_fill
        cell.font = name_font
        cell.alignment = Alignment(horizontal="left")
        for col, date in enumerate(sorted_dates, start=2):
            val = emp_data[name].get(date, "")
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = row_fill
            if val:
                cell.font = seat_font
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.font = Font(color="3B3760", name="Calibri")

    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="seating-{month}.xlsx"'}
    )


@router.get("/seating")
async def get_seating(month: str, date: str | None = None):
    async with get_db() as db:
        if date:
            rows = await (await db.execute(
                "SELECT employee_name, seat_id, status, is_manual, date"
                " FROM seat_assignments WHERE date=?",
                (date,)
            )).fetchall()
        else:
            rows = await (await db.execute(
                "SELECT employee_name, seat_id, status, is_manual, date"
                " FROM seat_assignments WHERE month=?",
                (month,)
            )).fetchall()

    assignments = [
        {"employee_name": r[0], "seat_id": r[1], "status": r[2],
         "is_manual": bool(r[3]), "date": r[4]}
        for r in rows
    ]
    return {"assignments": assignments}


@router.put("/seating/{date}/{seat_id}")
async def override_seat(date: str, seat_id: str, body: dict):
    employee_name = body.get("employee_name")
    month = date[:7]

    async with get_db() as db:
        if employee_name:
            await db.execute(
                "UPDATE seat_assignments SET seat_id=NULL WHERE date=? AND seat_id=? AND employee_name!=?",
                (date, seat_id, employee_name)
            )
            await db.execute(
                "UPDATE seat_assignments SET seat_id=NULL WHERE date=? AND employee_name=?",
                (date, employee_name)
            )
            await db.execute(
                """INSERT INTO seat_assignments(month,date,employee_name,seat_id,status,is_manual)
                   VALUES(?,?,?,?,'OFFICE',1)
                   ON CONFLICT(date,employee_name) DO UPDATE SET seat_id=?,status='OFFICE',is_manual=1""",
                (month, date, employee_name, seat_id, seat_id)
            )
        else:
            await db.execute(
                "UPDATE seat_assignments SET seat_id=NULL,is_manual=1 WHERE date=? AND seat_id=?",
                (date, seat_id)
            )
        await db.commit()
    return {"ok": True}